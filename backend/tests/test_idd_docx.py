from io import BytesIO
import re
import unittest
from types import SimpleNamespace

from docx import Document

from app.services.idd_docx import render_idd_docx


def node(name):
    return SimpleNamespace(name=name)


def field(name, *, size=2, unit="", value_range="", description="", purpose="", type="uint16"):
    return SimpleNamespace(
        id=1, name=name, type=type, type_kind="BASIC", ref_message=None,
        size_bytes=size, unit=unit, value_range=value_range,
        description=description, purpose=purpose, note="", order_index=1,
    )


def message(mid, name, *, tx=(), rx=(), protocol="TCP", fields=(), period="100",
            definition_type="STRUCT", order_index=0, size_bytes=2):
    return SimpleNamespace(
        id=mid, struct_name="MSG_" + str(mid), name=name,
        tx_targets=[node(x) for x in tx], rx_targets=[node(x) for x in rx],
        period=period, protocol=protocol, definition_type=definition_type,
        fields=list(fields), order_index=order_index, size_bytes=size_bytes,
    )


def document_text(doc):
    return "\n".join(
        [paragraph.text for paragraph in doc.paragraphs] +
        [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    )


class IddDocxTests(unittest.TestCase):
    def assert_sequential_idd_elements(self, doc, expected):
        """3.1 outward IE order and 3.2 detailed IE order must agree."""
        listing = next(table for table in doc.tables if table.rows[0].cells[0].text == "식별자")
        outward_ids = [re.match(r"IE-\d{4}", row.cells[6].text).group()
                       for row in listing.rows[1:]]
        self.assertEqual(outward_ids, expected[:len(outward_ids)])
        detailed = [p.text for p in doc.paragraphs
                    if re.match(r"^3\.2\.\d+\.\d+ IE-\d{4}\(", p.text)]
        detailed_ids = [re.search(r"IE-\d{4}", line).group() for line in detailed]
        self.assertEqual(detailed_ids, expected)
        if detailed:
            self.assertTrue(detailed[0].startswith("3.2.1.1 IE-0001("))


    def test_global_sequence_with_interleaved_relations_and_unsent_first(self):
        project = SimpleNamespace(name="TEST")
        messages = [
            message(10, "미연결", tx=(), rx=("CSU",), order_index=0),
            message(11, "항법", tx=("CSU",), rx=("AUV",), order_index=1),
            message(12, "상태", tx=("MDV",), rx=("CSU",), order_index=2),
            message(13, "임무", tx=("CSU",), rx=("AUV",), order_index=3),
            message(14, "회수", tx=("MDV",), rx=("CSU",), order_index=4),
        ]
        doc = Document(BytesIO(render_idd_docx(project, messages)))
        # A later message in an earlier-encountered relation must not jump IE IDs
        # when exported 3.2 chapters are grouped by relation.
        self.assert_sequential_idd_elements(
            doc, ["IE-0001", "IE-0002", "IE-0003", "IE-0004", "IE-0005"]
        )
        listing = next(table for table in doc.tables if table.rows[0].cells[0].text == "식별자")
        self.assertEqual(
            [row.cells[2].text for row in listing.rows[1:]],
            ["CSU", "CSU", "MDV", "MDV"],
        )
        self.assertEqual(
            [row.cells[6].text for row in listing.rows[1:]],
            [
                "IE-0001(항법(MSG_11))",
                "IE-0002(임무(MSG_13))",
                "IE-0003(상태(MSG_12))",
                "IE-0004(회수(MSG_14))",
            ],
        )
        details = [(table.rows[0].cells[2].text, table.rows[1].cells[2].text)
                   for table in doc.tables
                   if len(table.rows) > 3 and table.rows[0].cells[0].text == "인터페이스명(식별자)"]
        self.assertEqual(details[0], ("IE-0001(항법(MSG_11))", "CSU"))
        self.assertEqual(details[-1], ("IE-0005(미연결(MSG_10))", ""))


    def test_idd_docx_produces_editable_tables_with_saved_data(self):
        project = SimpleNamespace(name="CSU 연동")
        messages = [
            message(1, "항법정보", tx=("CSU",), rx=("AUV",), protocol="UDP, DDS",
                    fields=[field("latitude", size=4, unit="deg",
                                  value_range="-90~90", description="위도")],
                    size_bytes=4),
            message(2, "임무명령", tx=("CSU",), rx=("AUV",), protocol="DDS",
                    fields=[field("command")], order_index=2),
        ]
        content = render_idd_docx(project, messages)
        assert content[:2] == b"PK"
        doc = Document(BytesIO(content))
        text = document_text(doc)
        for value in ["3.1 인터페이스간 연동 개념", "3.2 인터페이스간 연동 설계",
                      "CSU", "AUV", "항법정보", "임무명령", "latitude",
                      "deg", "-90~90", "위도", "100 ms", "4 Byte", "UDP", "DDS",
                      "사용자 인터페이스 목록", "통신방법", "프로토콜"]:
            assert value in text
        assert doc.tables


    def test_idd_does_not_infer_missing_transport_details_or_pairwise_edges(self):
        project = SimpleNamespace(name="TEST")
        messages = [
            message(1, "A", tx=("CSU", "MDV"), rx=("AUV",), protocol="", fields=[]),
            message(2, "Only enum", definition_type="ENUM"),
        ]
        doc = Document(BytesIO(render_idd_docx(project, messages)))
        text = document_text(doc)
        assert "Only enum" not in text
        assert "CSU, MDV" in text
        # The optional TX/RX summary table is removed. Each sender gets one
        # outgoing interface row even when the same message has multiple senders.
        assert "송신 CSCI" not in [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        interface_table = next(table for table in doc.tables if table.rows[0].cells[0].text == "식별자")
        assert [row.cells[2].text for row in interface_table.rows[1:]] == ["CSU", "MDV"]
        assert [row.cells[6].text for row in interface_table.rows[1:]] == [
            "IE-0001(A(MSG_1))", "IE-0002(A(MSG_1))"
        ]
        assert all(row.cells[1].text == "" for row in interface_table.rows[1:])
        # Unknown communication fields stay blank, rather than fabricated.
        comm_table = next(
            table for table in doc.tables
            if table.rows[1].cells[0].text == "통신 구분명(식별자)"
        )
        assert comm_table.rows[1].cells[1].text == ""
        assert "3a" not in text


    def test_interface_list_uses_one_row_per_sender_and_sent_message(self):
        project = SimpleNamespace(name="TEST")
        messages = [
            message(1, "항법정보", tx=("CSU",), rx=("AUV",), protocol="UDP"),
            message(2, "임무명령", tx=("CSU", "MDV"), rx=("AUV",), protocol="DDS", order_index=1),
            message(3, "상태보고", tx=("AUV",), rx=("CSU",), protocol="TCP", order_index=2),
            message(4, "송신 미지정", tx=(), rx=("CSU",), order_index=3),
        ]
        doc = Document(BytesIO(render_idd_docx(project, messages)))
        interface_table = next(table for table in doc.tables if table.rows[0].cells[0].text == "식별자")
        assert len(interface_table.rows) == 5
        actual = [
            (row.cells[1].text, row.cells[2].text, row.cells[6].text)
            for row in interface_table.rows[1:]
        ]
        assert actual == [
            ("", "CSU", "IE-0001(항법정보(MSG_1))"),
            ("", "CSU", "IE-0002(임무명령(MSG_2))"),
            ("", "MDV", "IE-0003(임무명령(MSG_2))"),
            ("", "AUV", "IE-0004(상태보고(MSG_3))"),
        ]
        text = document_text(doc)
        assert "연동 노드 구성 (CSCI 송수신 관계)" not in text
        assert "IE-0005(송신 미지정(MSG_4))" in text  # 3.2 still documents it
        assert "IE-0002(임무명령(MSG_2))" in text
        assert "IE-0003(임무명령(MSG_2))" in text
        # Detail data-element rows and headings use the same IE identifier.
        assert "IE-0001(항법정보(MSG_1))" in text
        assert "IE-0004(상태보고(MSG_3))" in text
        self.assert_sequential_idd_elements(doc, ["IE-0001", "IE-0002", "IE-0003", "IE-0004", "IE-0005"])


    def test_idd_generates_template_when_project_is_empty(self):
        doc = Document(BytesIO(render_idd_docx(SimpleNamespace(name="Empty"), [])))
        assert doc.tables
        assert "등록된 메시지가 없습니다." in document_text(doc)
