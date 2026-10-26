"""Regression checks for HWPX IDD export.

Run from backend/: python -m unittest discover -s tests -v
These tests create and inspect real ZIP/OWPML bytes without Hancom Office.
"""
import unittest
import re
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZIP_STORED, ZipFile

from docx import Document
from lxml import etree

from app.services.idd_docx import render_idd_docx
from app.services.idd_hwpx import HP, HH, HWP_MIMETYPE, render_idd_hwpx


def node(name):
    return SimpleNamespace(name=name)


def field(name, description, **kwargs):
    return SimpleNamespace(
        id=1, name=name, type="uint16", type_kind="BASIC",
        ref_message=None, size_bytes=2,
        unit=kwargs.get("unit", ""), value_range=kwargs.get("value_range", ""),
        purpose="", description=description, note="", order_index=0,
    )


def message(mid, *, name="항법정보", tx=("CSU",), rx=("AUV",),
            protocol="UDP", fields=None):
    fields = fields if fields is not None else [
        field("latitude", "위도 <degree> & 각도", unit="deg", value_range="-90~90")
    ]
    return SimpleNamespace(
        id=mid, name=name, struct_name="MSG_%d" % mid,
        protocol=protocol, tx_targets=[node(x) for x in tx],
        rx_targets=[node(x) for x in rx], period="100", size_bytes=2,
        fields=fields, definition_type="STRUCT", order_index=0,
    )


class IddHwpxTests(unittest.TestCase):
    def test_hwpx_structure_and_all_editable_tables(self):
        project = SimpleNamespace(name="시험체계")
        messages = [message(1), message(2, name="명령상태", fields=[])]
        data = render_idd_hwpx(project, messages)
        word = Document(BytesIO(render_idd_docx(project, messages)))

        with ZipFile(BytesIO(data), "r") as archive:
            self.assertIsNone(archive.testzip())
            self.assertEqual("mimetype", archive.namelist()[0])
            self.assertEqual(HWP_MIMETYPE, archive.read("mimetype"))
            self.assertEqual(ZIP_STORED, archive.getinfo("mimetype").compress_type)
            for name in [
                "version.xml", "settings.xml",
                "Contents/header.xml", "Contents/content.hpf",
                "Contents/section0.xml", "META-INF/container.xml",
            ]:
                self.assertIn(name, archive.namelist())
            for name in archive.namelist():
                if name.endswith((".xml", ".hpf")):
                    etree.fromstring(archive.read(name))
            section = etree.fromstring(archive.read("Contents/section0.xml"))
            tables = section.findall(".//{%s}tbl" % HP)
            self.assertEqual(len(word.tables), len(tables))
            self.assertGreater(len(tables), 5)
            for hwpx_table, docx_table in zip(tables, word.tables):
                self.assertEqual(int(hwpx_table.get("colCnt")), len(docx_table.columns))
                self.assertEqual(int(hwpx_table.get("rowCnt")), len(docx_table.rows))
                rows = hwpx_table.findall("{%s}tr" % HP)
                for row in rows:
                    self.assertEqual(len(row.findall("{%s}tc" % HP)), len(docx_table.columns))
            section_text = "\n".join((t.text or "") for t in section.findall(".//{%s}t" % HP))
            for value in ("시험체계", "CSU", "AUV", "latitude",
                          "위도 <degree>", "& 각도", "명령상태", "-90~90",
                          "프로토콜", "통신방법", "사용자 인터페이스 목록"):
                self.assertIn(value, section_text)

    def test_hwpx_heading_ten_point_body_nine_point_and_renumbered_messages(self):
        project = SimpleNamespace(name="예제")
        messages = [
            message(1, name="정보보고", tx=("CSU", "MDV")),
            message(2, name="상태보고", tx=("CSU",)),
        ]
        with ZipFile(BytesIO(render_idd_hwpx(project, messages))) as archive:
            head = etree.fromstring(archive.read("Contents/header.xml"))
            table_header_style = head.find(".//{%s}charPr[@id='26']" % HH)
            table_body_style = head.find(".//{%s}charPr[@id='27']" % HH)
            self.assertEqual(table_header_style.get("height"), "1000")
            self.assertEqual(table_body_style.get("height"), "900")

            section = etree.fromstring(archive.read("Contents/section0.xml"))
            headings = section.findall(".//{%s}tbl/{%s}tr[1]//{%s}run" % (HP, HP, HP))
            bodies = section.findall(".//{%s}tbl/{%s}tr[2]//{%s}run" % (HP, HP, HP))
            self.assertTrue(headings and bodies)
            self.assertTrue(all(run.get("charPrIDRef") == "26" for run in headings))
            self.assertTrue(all(run.get("charPrIDRef") == "27" for run in bodies))
            values = [t.text or "" for t in section.findall(".//{%s}t" % HP)]
            text = " ".join(values)
            self.assertNotIn("연동 노드 구성 (CSCI 송수신 관계)", text)
            self.assertIn("IE-0001(정보보고(MSG_1))", text)
            self.assertIn("IE-0002(정보보고(MSG_1))", text)
            self.assertIn("IE-0003(상태보고(MSG_2))", text)

            interface_table = next(
                table for table in section.findall(".//{%s}tbl" % HP)
                if "식별자" in "".join(table.findall("{%s}tr" % HP)[0].itertext())
                and "연동기능" in "".join(table.findall("{%s}tr" % HP)[0].itertext())
            )
            records = []
            for row in interface_table.findall("{%s}tr" % HP)[1:]:
                cells = row.findall("{%s}tc" % HP)
                records.append(["".join(cell.itertext()) for cell in cells])
            self.assertEqual(len(records), 3)
            self.assertEqual([row[2] for row in records], ["CSU", "MDV", "CSU"])
            self.assertEqual([row[6] for row in records], [
                "IE-0001(정보보고(MSG_1))",
                "IE-0002(정보보고(MSG_1))",
                "IE-0003(상태보고(MSG_2))",
            ])
            self.assertTrue(all(row[1] == "" for row in records))
            detail_headings = []
            for paragraph in section.findall("{%s}p" % HP):
                label = "".join(
                    t.text or "" for t in paragraph.findall("./{%s}run/{%s}t" % (HP, HP))
                )
                if re.match(r"^3\.2\.\d+\.\d+ IE-\d{4}\(", label):
                    detail_headings.append(label)
            self.assertEqual(
                [re.search(r"IE-\d{4}", label).group() for label in detail_headings],
                ["IE-0001", "IE-0002", "IE-0003"],
            )
            self.assertTrue(detail_headings[0].startswith("3.2.1.1 IE-0001("))


    def test_missing_values_are_blank_and_user_input_is_xml_escaped(self):
        data = render_idd_hwpx(
            SimpleNamespace(name="A&B <C>"), [
                message(1, name="상태<&>정보", protocol="", tx=("CSU", "MDV"))
            ]
        )
        with ZipFile(BytesIO(data)) as archive:
            raw = archive.read("Contents/section0.xml")
            root = etree.fromstring(raw)
            text = " ".join(t.text or "" for t in root.findall(".//{%s}t" % HP))
            self.assertIn("상태<&>정보", text)
            self.assertIn("CSU, MDV", text)
            # The unknown communication category stays an editable empty cell.
            rows = root.findall(".//{%s}tr" % HP)
            self.assertTrue(any(
                "통신 구분명(식별자)" in "".join(tc.itertext())
                for row in rows for tc in row.findall("{%s}tc" % HP)
            ))
            self.assertIn(b"&lt;", raw)
            self.assertIn(b"&amp;", raw)

    def test_empty_project_has_structurally_valid_hwpx(self):
        data = render_idd_hwpx(SimpleNamespace(name="비어 있는 프로젝트"), [])
        with ZipFile(BytesIO(data)) as archive:
            root = etree.fromstring(archive.read("Contents/section0.xml"))
            self.assertTrue(root.findall("{%s}p" % HP))
            self.assertGreater(len(root.findall(".//{%s}tbl" % HP)), 0)


if __name__ == "__main__":
    unittest.main()
