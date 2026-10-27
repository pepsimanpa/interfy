"""Generate editable IDD chapter 3.1 and 3.2 tables from Interfy data.

The source guide is SW기술문서-04 (IDD), revised 2024-12-26.
Interfy integration targets are treated as CSCIs. Since this application does
not classify relationships into internal/external interfaces, corresponding
guide tables are combined, not duplicated or assigned a guessed category.
Unavailable attributes are deliberately exported as empty Word cells.
"""
from __future__ import annotations

from collections import OrderedDict
from io import BytesIO
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _name(target) -> str:
    return _text(getattr(target, "name", ""))


def _unique(values):
    return list(dict.fromkeys(value for value in values if value))


def _target_names(message, kind):
    targets = getattr(message, "%s_targets" % kind, None) or []
    return _unique(_name(target) for target in targets)


def _protocol_names(messages):
    result = []
    for message in messages:
        for part in re.split(r"[,;|+\n\r]|\s+/\s+", _text(getattr(message, "protocol", ""))):
            part = part.strip()
            if part and part not in result:
                result.append(part)
    return result


def _message_display(message):
    return _text(getattr(message, "name", "")) or _text(getattr(message, "struct_name", ""))


def _message_identity(message):
    label = _message_display(message)
    technical = _text(getattr(message, "struct_name", ""))
    if technical and technical != label:
        return "%s (%s)" % (label, technical)
    return label


def _relationships(messages):
    """Group by the *whole* TX/RX sets without inventing all-to-all edges."""
    result = OrderedDict()
    ordered = sorted(messages, key=lambda m: (getattr(m, "order_index", 0) or 0, getattr(m, "id", 0) or 0))
    for message in ordered:
        if _text(getattr(message, "definition_type", "STRUCT")).upper() == "ENUM":
            continue
        tx = tuple(sorted(_target_names(message, "tx")))
        rx = tuple(sorted(_target_names(message, "rx")))
        result.setdefault((tx, rx), []).append(message)
    return result


def _sequenced_data_elements(relations):
    """Build the shared 3.1/3.2 order and one document-wide IE counter.

    One IE represents a specific sending node's message occurrence. The same
    message sent by two different CSCIs is therefore assigned *two* IE IDs.
    Within each relation group, group messages by sender to keep 3.1 readable.
    Unassigned-sender messages are documented in 3.2 after the outgoing items
    and do not produce a fictitious 3.1 interface entry.
    """
    result = OrderedDict()
    next_number = 1
    # Stable sort: assigned-sender groups first so the very first 3.2 data
    # element is IE-0001 and 3.1 also starts with IE-0001.
    keys = sorted(relations, key=lambda key: not bool(key[0]))
    for key in keys:
        senders, _receivers = key
        group = relations[key]
        entries = []
        if senders:
            for sender in senders:
                for message in group:
                    entries.append((sender, message, "IE-%04d" % next_number))
                    next_number += 1
        else:
            for message in group:
                entries.append(("", message, "IE-%04d" % next_number))
                next_number += 1
        result[key] = entries
    return result


def _data_element_label(message, identifier):
    name = _message_display(message)
    struct = _text(getattr(message, "struct_name", ""))
    detail = "%s(%s)" % (name, struct) if struct else name
    return "%s(%s)" % (identifier, detail)


def _period(message):
    value = _text(getattr(message, "period", ""))
    if not value or value == "0":
        return "비주기"
    if value.isdigit():
        return "%s ms" % value
    return value


def _size(value):
    try:
        size = getattr(value, "size_bytes", None)
        return str(size) if size is not None and size != 0 else ""
    except (RecursionError, ValueError, TypeError, AttributeError):
        return ""


def _shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    tc_pr.append(shading)


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def _table(document, columns, rows, widths=None):
    """Use actual editable Word tables. No unknown value is replaced by '-'."""
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.autofit = False
    available_cm = 17.0
    widths = widths or [available_cm / len(columns)] * len(columns)
    for index, heading in enumerate(columns):
        table.columns[index].width = Cm(widths[index])
        cell = table.rows[0].cells[index]
        cell.width = Cm(widths[index])
        cell.text = heading
        _shade(cell, "E8EDF3")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(8)
    _repeat_header(table.rows[0])
    for values in rows or [[""] * len(columns)]:
        cells = table.add_row().cells
        for index in range(len(columns)):
            cells[index].width = Cm(widths[index])
            cells[index].text = _text(values[index] if index < len(values) else "")
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8)
    document.add_paragraph()
    return table


def _caption(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = False
    run.font.name = "맑은 고딕"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0, 0, 0)


def _heading(document, text, level):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def _field_type(field):
    kind = _text(getattr(field, "type_kind", "BASIC")).upper()
    if kind in ("MESSAGE", "ENUM"):
        ref = getattr(field, "ref_message", None)
        if ref is not None:
            return _text(getattr(ref, "struct_name", "")) or _text(getattr(ref, "name", ""))
    return _text(getattr(field, "type", ""))


def _field_detail(field):
    return "\n".join(_unique([
        _text(getattr(field, "purpose", "")),
        _text(getattr(field, "description", "")),
        _text(getattr(field, "note", "")),
    ]))


def _data_element_table(document, message, identifier, sender):
    """Use one page-width IDD table with the same Word styles as other tables.

    Setting grid-column widths alone does not resize python-docx's
    pre-existing cell widths. In particular, horizontal cell merges can then
    add up the *default* equal-width cells and grow past the page margins.
    Set every cell width before merging, then apply merged widths explicitly.
    """
    element = _data_element_label(message, identifier)
    _caption(document, "데이터요소 - " + element)
    columns = ["요소명", "타입", "크기", "단위", "보안", "범위", "요소내용"]
    # The surrounding document is A4 with 2cm margins (17cm usable width).
    widths = [2.0, 1.35, 1.0, 1.15, 1.05, 2.35, 8.1]
    table = document.add_table(rows=4, cols=len(columns))
    table.style = "Table Grid"
    table.autofit = False
    for idx, width in enumerate(widths):
        table.columns[idx].width = Cm(width)

    def set_row_widths(row):
        for index, cell in enumerate(row.cells):
            cell.width = Cm(widths[index])

    for row in table.rows:
        set_row_widths(row)

    def merged_value(row_index, first, last, value):
        # All underlying cells already have the correct widths.
        merged = table.cell(row_index, first).merge(table.cell(row_index, last))
        merged.width = Cm(sum(widths[first:last + 1]))
        merged.text = _text(value)
        return merged

    merged_value(0, 0, 1, "인터페이스명(식별자)")
    merged_value(0, 2, 6, element)
    merged_value(1, 0, 1, "송신처")
    merged_value(1, 2, 3, sender)
    merged_value(1, 4, 5, "수신처")
    table.cell(1, 6).text = ", ".join(_target_names(message, "rx"))
    merged_value(2, 0, 1, "발생주기")
    merged_value(2, 2, 3, _period(message))
    merged_value(2, 4, 5, "데이터 길이")
    table.cell(2, 6).text = ("%s Byte" % _size(message)) if _size(message) else ""

    for index, heading in enumerate(columns):
        table.cell(3, index).text = heading

    fields = sorted(getattr(message, "fields", []) or [],
                    key=lambda field: (getattr(field, "order_index", 0) or 0,
                                       getattr(field, "id", 0) or 0))
    for field in fields:
        row = table.add_row()
        set_row_widths(row)
        values = [_text(getattr(field, "name", "")), _field_type(field), _size(field),
                  _text(getattr(field, "unit", "")), "",
                  _text(getattr(field, "value_range", "")), _field_detail(field)]
        for index, value in enumerate(values):
            row.cells[index].text = value

    # Metadata label cells and the field-header row reuse the same light-blue
    # fill, bold 8pt type and alignment used by the ordinary Word tables.
    for row_index, row in enumerate(table.rows):
        seen = set()
        for index, cell in enumerate(row.cells):
            if id(cell._tc) in seen:
                continue  # merged cells appear multiple times in row.cells
            seen.add(id(cell._tc))
            label = (row_index == 3 or
                     (row_index == 0 and index == 0) or
                     (row_index in (1, 2) and index in (0, 4)))
            if label:
                _shade(cell, "E8EDF3")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                if label:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "맑은 고딕"
                    run.font.size = Pt(8)
                    run.bold = bool(label)
    _repeat_header(table.rows[3])
    document.add_paragraph()


def _communication_table(document, tx, rx):
    _caption(document, "통신방법 (표 9·12 통합)")
    _table(document, ["항목", "내용", "항목", "내용"], [
        ["통신 구분명(식별자)", "", "연동체계", ", ".join(rx)],
        ["연동 Type", "", "연동망", ""],
        ["통신매체 및 대역폭", "", "", ""],
        ["빈도 및 특성", "", "", ""],
        ["메시지 형식", "", "", ""],
        ["흐름제어", "", "", ""],
        ["데이터 전송률/전송간격", "", "", ""],
        ["어드레싱, 라우팅 규칙", "", "", ""],
        ["우선순위, 등급", "", "", ""],
        ["안전/보안 등 고려사항", "", "", ""],
    ], [4.0, 4.5, 4.0, 4.5])


def _protocol_table(document, protocol, rx):
    _caption(document, "프로토콜 (표 10 형식)")
    _table(document, ["항목", "내용", "항목", "내용"], [
        ["프로토콜명(식별자)", protocol, "연동체계", ", ".join(rx)],
        ["계층", "", "", ""],
        ["패킷", "", "", ""],
        ["에러제어 및 복구", "", "", ""],
        ["동기화", "", "", ""],
        ["상태 및 기타 특성", "", "", ""],
    ], [4.0, 4.5, 4.0, 4.5])


def render_idd_docx(project, messages) -> bytes:
    """Return an in-memory editable Word export for IDD chapters 3.1 and 3.2."""
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    normal = document.styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(25, 35, 50)
    normal.paragraph_format.space_after = Pt(5)

    _heading(document, "%s 인터페이스설계기술서(IDD)" % _text(getattr(project, "name", "")), 0)
    document.add_paragraph("3.1~3.2 인터페이스 설계 (Interfy 자동 작성)")
    document.add_paragraph("Interfy 노드는 CSCI로 취급하며 내부·외부 구분 없이 연동관계별로 정리합니다. 수집되지 않은 항목은 빈칸입니다.")

    relations = _relationships(messages)
    sequenced_relations = _sequenced_data_elements(relations)
    outgoing_entries = [
        (sender, message, identifier)
        for entries in sequenced_relations.values()
        for sender, message, identifier in entries
        if sender
    ]

    _heading(document, "3.1 인터페이스간 연동 개념", 1)
    _heading(document, "3.1.1 인터페이스 대상 식별과 다이어그램", 2)
    # No extra CSCI TX/RX relationship summary is needed in the exported IDD.
    _heading(document, "3.1.1.1 인터페이스 목록", 3)
    _caption(document, "표 5·6 통합 인터페이스 목록")
    _table(document,
           ["식별자", "연동기능", "연동대상", "통신방법", "상호운용성 수준", "프로토콜", "데이터요소"],
           [["", "", sender, "", "", ", ".join(_protocol_names([message])),
             _data_element_label(message, identifier)]
            for sender, message, identifier in outgoing_entries],
           [1.3, 1.8, 2.3, 1.8, 1.2, 1.7, 6.9])
    _heading(document, "3.2 인터페이스간 연동 설계", 1)
    if not sequenced_relations:
        document.add_paragraph("등록된 메시지가 없습니다.")
    for number, ((tx, rx), entries) in enumerate(sequenced_relations.items(), start=1):
        left = ", ".join(tx)
        right = ", ".join(rx)
        relation_name = " → ".join(part for part in (left, right) if part) or "연동 노드 미지정"
        _heading(document, "3.2.%d %s" % (number, relation_name), 2)

        # This export intentionally contains data-element details only.
        # Do not append communication, protocol or physical-compatibility
        # sections, or their now-unused placeholder tables.
        for sub, (sender, message, identifier) in enumerate(entries, start=1):
            element = _data_element_label(message, identifier)
            _heading(document, "3.2.%d.%d %s 데이터요소" % (number, sub, element), 3)
            _data_element_table(document, message, identifier, sender)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
