"""Export Interfy's existing IDD Word report as native, editable HWPX (OWPML).

We intentionally reuse render_idd_docx() as the single source of document
content, instead of duplicating CSCI/message/field mappings.

The HWPX skeleton, report header, and table cell templates are MIT-licensed
assets adapted from https://github.com/kist-aix/hwpx (see hwpx_assets/LICENSE).
HWPX is a ZIP of XML files. No Hancom Office, LibreOffice, or OS automation
is required on the backend.
"""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from itertools import count
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

from docx import Document
from docx.oxml.ns import qn as word_qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree

from .idd_docx import render_idd_docx

ASSETS = Path(__file__).resolve().parent / "hwpx_assets"
HP = "http://www.hancom.co.kr/hwpml/2011/paragraph"
HH = "http://www.hancom.co.kr/hwpml/2011/head"
OPF = "http://www.idpf.org/2007/opf/"
HWP_MIMETYPE = b"application/hwp+zip"

PACKAGE_FILES = (
    "version.xml",
    "settings.xml",
    "META-INF/container.xml",
    "META-INF/container.rdf",
    "META-INF/manifest.xml",
    "Contents/content.hpf",
    "Contents/header.xml",
    "Contents/section0.xml",
    "Preview/PrvText.txt",
)

_XML_PARSER = etree.XMLParser(
    resolve_entities=False, load_dtd=False, no_network=True, huge_tree=False
)


def _asset_xml(path):
    return etree.parse(str(ASSETS / path), parser=_XML_PARSER)


def _set_run_text(run, value):
    """Preserve line breaks as editable HWPX text instead of literal newlines."""
    for text_node in list(run.findall("{%s}t" % HP)):
        run.remove(text_node)
    parts = str(value or "").split("\n")
    for index, part in enumerate(parts):
        if index:
            etree.SubElement(run, "{%s}lineBreak" % HP)
        node = etree.SubElement(run, "{%s}t" % HP)
        node.text = part or None


def _paragraph(text, ids, *, heading=False, title=False, caption=False):
    # Every non-table paragraph except the report title is explicitly 11pt.
    char_style = "10" if title else ("29" if caption else ("28" if heading else "25"))
    p = etree.Element("{%s}p" % HP, {
        "id": str(next(ids)),
        "paraPrIDRef": "18" if caption else ("2" if heading or title else "22"),
        "styleIDRef": "0" if heading or title else "1",
        "pageBreak": "0",
        "columnBreak": "0",
        "merged": "0",
    })
    run = etree.SubElement(p, "{%s}run" % HP, {
        "charPrIDRef": char_style,
    })
    _set_run_text(run, text)
    return p


def _column_widths(table, available=42520):
    """Scale source Word column widths to the HWPX A4 printable width."""
    if not table.rows:
        return []
    raw = [
        int(column.width or 0)
        for column in table.columns
    ]
    if not any(raw):
        raw = [1] * len(raw)
    raw = [max(1, x) for x in raw]
    total = sum(raw)
    widths = [max(1, round(available * width / total)) for width in raw]
    widths[-1] += available - sum(widths)
    return widths


def _fill_cell(tc, text, row, col, width, ids):
    addr = tc.find("{%s}cellAddr" % HP)
    addr.set("rowAddr", str(row))
    addr.set("colAddr", str(col))
    cell_sz = tc.find("{%s}cellSz" % HP)
    cell_sz.set("width", str(width))
    paragraphs = tc.findall(".//{%s}p" % HP)
    for p in paragraphs:
        p.set("id", str(next(ids)))
    run = tc.find(".//{%s}run" % HP)
    _set_run_text(run, text)


def _table(table, ids, table_template, caption_text=None):
    """Create actual hp:tbl / hp:tr / hp:tc objects, not tab-delimited text."""
    if not table.rows:
        return None

    layout = _column_widths(table)
    width_total = sum(layout)
    source = table_template.find(".//{%s}tbl" % HP)
    head_row = source.find("{%s}tr" % HP)
    body_row = source.findall("{%s}tr" % HP)[1]
    head_cell = head_row.find("{%s}tc" % HP)
    body_cell = body_row.find("{%s}tc" % HP)

    tbl = deepcopy(source)
    tbl.set("id", str(next(ids)))
    tbl.set("rowCnt", str(len(table.rows)))
    tbl.set("colCnt", str(len(layout)))
    tbl.set("noAdjust", "0")
    tbl.find("{%s}pos" % HP).set("treatAsChar", "0")
    shape = tbl.find("{%s}sz" % HP)
    shape.set("width", str(width_total))
    shape.set("height", str(2363 + max(0, len(table.rows) - 1) * 1984))
    for tr in list(tbl.findall("{%s}tr" % HP)):
        tbl.remove(tr)

    # A merged Word cell occupies several grid positions but only one HWPX cell.
    is_data_element = len(table.rows) > 3 and table.rows[0].cells[0].text == "인터페이스명(식별자)"
    for row_index, docx_row in enumerate(table.rows):
        is_header = (row_index == 3) if is_data_element else (row_index == 0)
        tr = deepcopy(head_row if is_header else body_row)
        tr.attrib.pop("row-type", None)
        for existing in list(tr.findall("{%s}tc" % HP)):
            tr.remove(existing)
        col_index = 0
        while col_index < len(layout):
            cell = docx_row.cells[col_index]
            span = 1
            while (col_index + span < len(layout)
                   and docx_row.cells[col_index + span]._tc is cell._tc):
                span += 1
            # In the reference IDD, the first three rows shade only the
            # label cells; field headings (row 4) are shaded throughout.
            is_label = is_header or (
                is_data_element and
                ((row_index == 0 and col_index == 0) or
                 (row_index in (1, 2) and col_index in (0, 4)))
            )
            tc = deepcopy(head_cell if is_label else body_cell)
            combined_width = sum(layout[col_index:col_index + span])
            _fill_cell(tc, "\n".join(p.text for p in cell.paragraphs),
                       row_index, col_index, combined_width, ids)
            tc.find("{%s}cellSpan" % HP).set("colSpan", str(span))
            tr.append(tc)
            col_index += span
        tbl.append(tr)

    if caption_text:
        caption = etree.Element("{%s}caption" % HP, {
            "side": "TOP", "fullSz": "0", "width": str(width_total),
            "gap": "566", "lastWidth": str(width_total),
        })
        sublist = deepcopy(body_cell.find("{%s}subList" % HP))
        for old in list(sublist):
            sublist.remove(old)
        caption_p = _paragraph(caption_text, ids, caption=True)
        sublist.append(caption_p)
        caption.append(sublist)
        # In OWPML the caption follows outMargin and precedes inMargin.
        out_margin = tbl.find("{%s}outMargin" % HP)
        tbl.insert(list(tbl).index(out_margin) + 1, caption)

    # A table is a control embedded in a paragraph/run in HWPX.
    p = etree.Element("{%s}p" % HP, {
        "id": str(next(ids)), "paraPrIDRef": "2", "styleIDRef": "0",
        "pageBreak": "0", "columnBreak": "0", "merged": "0",
    })
    run = etree.SubElement(p, "{%s}run" % HP, {"charPrIDRef": "27"})
    run.append(tbl)
    etree.SubElement(run, "{%s}t" % HP)
    return p


def _hwpx_section(docx):
    tree = _asset_xml("Contents/section0.xml")
    root = tree.getroot()
    ids = count(1000000100)
    basic = _asset_xml("basic-table.xml").getroot()

    pending_caption = None
    for child in docx.element.body.iterchildren():
        if child.tag == word_qn("w:p"):
            paragraph = Paragraph(child, docx)
            if paragraph.text.startswith("데이터요소 - "):
                pending_caption = paragraph.text
                continue
            if paragraph.text.startswith("표 5·6 "):
                # Caption for the summary interface list: centered, black 11pt.
                root.append(_paragraph(paragraph.text, ids, caption=True))
                continue
            style_name = paragraph.style.name if paragraph.style is not None else ""
            is_title = style_name.startswith("Title")
            is_heading = style_name.startswith("Heading")
            root.append(_paragraph(paragraph.text, ids, heading=is_heading, title=is_title))
        elif child.tag == word_qn("w:tbl"):
            table = Table(child, docx)
            element = _table(table, ids, basic, caption_text=pending_caption)
            pending_caption = None
            if element is not None:
                root.append(element)

    return etree.tostring(tree, encoding="UTF-8", xml_declaration=True)


def _hwpx_header():
    """Use explicit HWPX character shapes that match the IDD report styles.

    Font sizes are hundredths of a point (1100=11pt). Use a real, widely
    available Korean Gothic for normal paragraphs rather than '고딕', which
    may be an uninstalled family name. The source PDF's table fonts are
    Dotum (labels) and Batang (values), approximately 10pt and 9pt.
    """
    tree = _asset_xml("Contents/header.xml")
    root = tree.getroot()
    fontfaces = root.find(".//{%s}fontfaces" % HH)
    langs = ("HANGUL", "LATIN", "HANJA", "JAPANESE", "OTHER", "SYMBOL", "USER")

    def set_face(char_pr, family):
        ref = char_pr.find("{%s}fontRef" % HH)
        if ref is None:
            raise ValueError("Missing HWPX fontRef")
        for lang in langs:
            group = fontfaces.find("{%s}fontface[@lang='%s']" % (HH, lang))
            existing = group.findall("{%s}font" % HH)
            entry = next((font for font in existing if font.get("face") == family), None)
            if entry is None:
                entry = etree.SubElement(group, "{%s}font" % HH, {
                    "id": str(len(existing)), "face": family, "type": "TTF",
                    "isEmbedded": "0",
                })
                group.set("fontCnt", str(len(existing) + 1))
            ref.set(lang.lower(), entry.get("id"))

    char_props = root.find(".//{%s}charProperties" % HH)
    base = char_props.find("{%s}charPr[@id='25']" % HH)
    if base is None:
        raise ValueError("Missing base HWPX body character style")
    for style_id, family, size in (
        ("25", "맑은 고딕", 11),
        ("26", "돋움", 10),
        ("27", "바탕", 9),
    ):
        char_pr = char_props.find("{%s}charPr[@id='%s']" % (HH, style_id))
        if char_pr is None:
            raise ValueError("Missing HWPX character style " + style_id)
        char_pr.set("height", str(size * 100))
        set_face(char_pr, family)
        # The source table has plain, not bold, Dotum column labels.
        if style_id == "26":
            bold = char_pr.find("{%s}bold" % HH)
            if bold is not None:
                char_pr.remove(bold)

    for style_id, family, is_bold in (
        ("28", "맑은 고딕", True),
        ("29", "맑은 고딕", False),
    ):
        new_style = deepcopy(base)
        new_style.set("id", style_id)
        new_style.set("height", "1100")
        new_style.set("textColor", "#000000" if style_id == "29" else "#0000FF")
        set_face(new_style, family)
        if is_bold:
            etree.SubElement(new_style, "{%s}bold" % HH)
        char_props.append(new_style)
    char_props.set("itemCnt", str(len(char_props.findall("{%s}charPr" % HH))))

    for style_id in ("0", "1"):
        style = root.find(".//{%s}style[@id='%s']" % (HH, style_id))
        if style is None:
            raise ValueError("Missing HWPX normal/body style " + style_id)
        style.set("charPrIDRef", "25")

    # Reference PDF page 14 fills metadata labels and field headings #D6D6D6.
    # HWPX cell templates use borderFill 9 for shaded and 10 for plain cells.
    gray_border = root.find(".//{%s}borderFill[@id='9']" % HH)
    if gray_border is None:
        raise ValueError("Missing table label borderFill 9")
    fill = gray_border.find(".//{http://www.hancom.co.kr/hwpml/2011/core}winBrush")
    if fill is None:
        raise ValueError("Missing table label fill brush")
    fill.set("faceColor", "#D6D6D6")
    return etree.tostring(tree, encoding="UTF-8", xml_declaration=True)


def _manifest(project_name):
    tree = _asset_xml("Contents/content.hpf")
    root = tree.getroot()
    title = root.find(".//{%s}title" % OPF)
    if title is not None:
        title.text = str(project_name or "Interfy") + " 인터페이스설계기술서(IDD)"
    return etree.tostring(tree, encoding="UTF-8", xml_declaration=True)


def render_idd_hwpx(project, messages) -> bytes:
    """Generate a Korean-editable HWPX with the same 3.1/3.2 tables as Word."""
    word_bytes = render_idd_docx(project, messages)
    word_doc = Document(BytesIO(word_bytes))
    section_xml = _hwpx_section(word_doc)
    header_xml = _hwpx_header()
    manifest_xml = _manifest(getattr(project, "name", ""))

    preview = "\n".join(
        paragraph.text for paragraph in word_doc.paragraphs if paragraph.text
    )[:2048].encode("utf-8")

    result = BytesIO()
    with ZipFile(result, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", HWP_MIMETYPE, compress_type=ZIP_STORED)
        for path in PACKAGE_FILES:
            if path == "Contents/section0.xml":
                data = section_xml
            elif path == "Contents/header.xml":
                data = header_xml
            elif path == "Contents/content.hpf":
                data = manifest_xml
            elif path == "Preview/PrvText.txt":
                data = preview
            else:
                data = (ASSETS / path).read_bytes()
            archive.writestr(path, data)

    return result.getvalue()
